"""Dataset persistence for a source: products.json, products.xlsx, reports.

Excel/JSON is the source dataset (source of truth for the demo). Azure AI Search
is the retrieval/index layer built from it.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from ..config import settings
from ..models.product import Product

PIPELINE_VERSION = "1.0.0"


def _source_dir(source_id: str) -> Path:
    d = settings.source_data_dir(source_id)
    d.mkdir(parents=True, exist_ok=True)
    return d


def products_json_path(source_id: str) -> Path:
    return _source_dir(source_id) / "products.json"


def products_xlsx_path(source_id: str) -> Path:
    return _source_dir(source_id) / "products.xlsx"


def embeddings_path(source_id: str) -> Path:
    return _source_dir(source_id) / "embeddings.json"


def ingestion_report_path(source_id: str) -> Path:
    return _source_dir(source_id) / "ingestion_report.json"


# ---------------------------------------------------------------- JSON dataset
def write_products_json(source_id: str, products: list[Product]) -> Path:
    path = products_json_path(source_id)
    payload = [p.model_dump() for p in products]
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def read_products_json(source_id: str) -> list[Product]:
    path = products_json_path(source_id)
    if not path.exists():
        return []
    data = json.loads(path.read_text(encoding="utf-8"))
    return [Product(**item) for item in data]


# --------------------------------------------------------------- Excel dataset
def _write_sheet(ws, headers: list[str], rows: list[dict[str, Any]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])


def write_products_xlsx(source_id: str, products: list[Product], metadata: dict[str, Any]) -> Path:
    path = products_xlsx_path(source_id)

    products_rows = [
        {
            "product_id": p.product_id,
            "source_id": p.source_id,
            "source_name": p.source_name,
            "brand_name": p.brand_name,
            "composition": p.composition,
            "therapeutic_category": p.therapeutic_category,
            "source_url": p.source_url,
            "source_title": p.source_title,
            "scraped_at": p.scraped_at,
        }
        for p in products
    ]

    ingredient_rows = [
        {
            "product_id": ing.product_id,
            "source_id": ing.source_id,
            "brand_name": ing.brand_name,
            "ingredient": ing.ingredient,
            "strength": ing.strength,
            "unit": ing.unit,
            "raw_composition": ing.raw_composition,
        }
        for p in products
        for ing in p.ingredients
    ]

    cat_map: dict[str, list[str]] = {}
    for p in products:
        cat = p.therapeutic_category or "(uncategorized)"
        cat_map.setdefault(cat, []).append(p.brand_name)
    category_rows = [
        {
            "source_id": source_id,
            "therapeutic_category": cat,
            "product_count": len(brands),
            "products": ", ".join(sorted(brands)),
        }
        for cat, brands in sorted(cat_map.items())
    ]

    wb = Workbook()
    ws_products = wb.active
    ws_products.title = "Products"
    _write_sheet(
        ws_products,
        ["product_id", "source_id", "source_name", "brand_name", "composition",
         "therapeutic_category", "source_url", "source_title", "scraped_at"],
        products_rows,
    )
    _write_sheet(
        wb.create_sheet("Ingredients"),
        ["product_id", "source_id", "brand_name", "ingredient", "strength", "unit", "raw_composition"],
        ingredient_rows,
    )
    _write_sheet(
        wb.create_sheet("Categories"),
        ["source_id", "therapeutic_category", "product_count", "products"],
        category_rows,
    )
    _write_sheet(
        wb.create_sheet("Metadata"),
        list(metadata.keys()),
        [metadata],
    )
    wb.save(path)
    return path


# ------------------------------------------------------------------- reports
def write_ingestion_report(source_id: str, report: dict[str, Any]) -> Path:
    path = ingestion_report_path(source_id)
    path.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def read_ingestion_report(source_id: str) -> dict[str, Any] | None:
    path = ingestion_report_path(source_id)
    if not path.exists():
        return None
    return json.loads(path.read_text(encoding="utf-8"))


# ---------------------------------------------------------------- embeddings
def write_embeddings(source_id: str, records: list[dict[str, Any]]) -> Path:
    path = embeddings_path(source_id)
    path.write_text(json.dumps(records, indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def read_embeddings(source_id: str) -> dict[str, dict[str, Any]]:
    """Return embeddings keyed by product_id."""
    path = embeddings_path(source_id)
    if not path.exists():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    return {rec["product_id"]: rec for rec in data}
